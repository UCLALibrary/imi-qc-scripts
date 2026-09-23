"""Filename-convention checker for UCLA Special Collections imaging manifests.
See README.md for the naming rules, the check list, and design rationale.

Usage:
    python manifest_qc_v2.py MANIFEST.xlsx [MORE.xlsx ...] [--status status.csv] [--out PATH]
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

FILENAME_RE = re.compile(
    r"^uclalsc_(?P<coll>[^_]+)_(?P<item>ms[^_]+)_(?P<seq>\d+)(?:_(?P<pos>.+))?\.tif$"
)
POS_BODY_RE = re.compile(r"^(?P<seg>[a-z]+)_?(?P<num>\d+)(?P<side>[rv])?$")

POSITION_SPECS = {
    "a": re.compile(r"^a_(\d+)$"),
    "b": re.compile(r"^b_(\d+)[rv]$"),
    "f": re.compile(r"^f_(\d+)[rv]$"),
    "y": re.compile(r"^y_(\d+)[rv]$"),
    "z": re.compile(r"^z_(\d+)$"),
    "frag": re.compile(r"^frag_?(\d+)[rv]$"),
    "p": re.compile(r"^p_(\d+)$"),
}

# frag with a segment letter glued in front (b_frag1r) is a different violation
# from frag's own optional underscore — see check_grammar / check_frag_style.
PREFIXED_FRAG_RE = re.compile(r"^[a-z]+_frag")

WHITELIST_SEGMENTS = frozenset(POSITION_SPECS)
SIDED_SEGMENTS = ("b", "f", "y", "frag")
COUNTER_SEGMENTS = ("a", "z", "p")
FOLIO_DEFAULT = "f"

OUTPUT_DIR = Path(__file__).resolve().parent / "output"
SEGMENT_RANK = {"a": 0, "b": 1, "f": 2, "p": 2, "y": 3, "z": 4}
VALID_STATUSES = ("foliated", "unfoliated", "paginated")
FLYLEAF_REVIEW_THRESHOLD = 4  # flag for a look past this many leaves; not a hard rule


@dataclass
class Position:
    segment: str
    number: int
    side: str | None
    raw_number: str
    family: str | None
    canonical: bool


@dataclass
class ParsedFile:
    name: str
    item: str
    seq: int
    pos: str | None
    position: Position | None


@dataclass
class Finding:
    item: str
    ref: str
    severity: str
    code: str
    message: str


def _family(segment: str, side: str | None) -> str | None:
    # an unrecognized segment with an r/v side still reads as a folio, so a
    # segment typo and a real ordering/completeness problem get separate findings
    if segment in WHITELIST_SEGMENTS:
        return segment
    return FOLIO_DEFAULT if side else None


def parse_position(pos: str) -> Position | None:
    body = POS_BODY_RE.match(pos)
    if not body:
        return None
    segment, raw_number, side = body.group("seg"), body.group("num"), body.group("side")
    canonical = any(spec.match(pos) for spec in POSITION_SPECS.values())
    return Position(segment, int(raw_number), side, raw_number,
                    _family(segment, side), canonical)


def parse_filename(name: str) -> ParsedFile | None:
    match = FILENAME_RE.match(name)
    if not match:
        return None
    pos = match.group("pos")
    position = parse_position(pos) if pos else None
    # collection+item: item numbers repeat across collections
    item_key = f"{match.group('coll')}_{match.group('item')}"
    return ParsedFile(name, item_key, int(match.group("seq")), pos, position)


def load_manifest(path: Path) -> list[str]:
    if path.suffix.lower() in {".txt", ".csv"}:
        return _load_text(path)
    return _load_xlsx(path)


def _load_text(path: Path) -> list[str]:
    firsts = (line.split(",")[0].strip().strip('"')
              for line in path.read_text(encoding="utf-8-sig").splitlines())
    return [f for f in firsts if f.lower().endswith(".tif")]


def _load_xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    # tab order varies between deliveries — take whichever sheet actually holds .tif names
    workbook = load_workbook(path, read_only=True, data_only=True)
    best: list[str] = []
    for sheet_name in workbook.sheetnames:
        names = [
            row[0].strip()
            for row in workbook[sheet_name].iter_rows(values_only=True)
            if row and isinstance(row[0], str) and row[0].strip().lower().endswith(".tif")
        ]
        if len(names) > len(best):
            best = names
    workbook.close()
    return best


def load_status(path: Path) -> dict[str, str]:
    status = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            if len(row) < 2 or row[0].strip().lower() == "item":
                continue
            item, value = row[0].strip(), row[1].strip().lower()
            if value not in VALID_STATUSES:
                raise ValueError(f"{item}: status '{value}' not one of {VALID_STATUSES}")
            status[item] = value
    return status


def check_parsing(item: str, names: list[str], parsed: list[ParsedFile]) -> list[Finding]:
    unparsed = set(names) - {p.name for p in parsed}
    return [
        Finding(item, name, "ERROR", "unparseable", "does not fit the naming pattern")
        for name in sorted(unparsed)
    ]


def check_grammar(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    groups = defaultdict(list)
    prefixed_frag = []
    for p in parsed:
        if p.pos is None:
            continue
        if PREFIXED_FRAG_RE.match(p.pos):
            prefixed_frag.append((p.seq, p.pos))
        elif p.position is None or not p.position.canonical:
            key = p.position.segment if p.position else _leading_alpha(p.pos)
            groups[key].append((p.seq, p.pos))
    findings = []
    if prefixed_frag:
        seqs = sorted(seq for seq, _ in prefixed_frag)
        findings.append(
            Finding(item, _compact(seqs), "ERROR", "frag-prefixed",
                    f"{len(prefixed_frag)} file(s) have a segment letter glued onto frag, "
                    f"e.g. '{prefixed_frag[0][1]}' — frag should never carry a leading segment code")
        )
    for key, entries in groups.items():
        seqs = sorted(seq for seq, _ in entries)
        findings.append(
            Finding(item, _compact(seqs), "ERROR", "naming",
                    _grammar_hint(key, entries[0][1], len(entries)))
        )
    return findings


def _leading_alpha(pos: str) -> str:
    match = re.match(r"^[a-z]+", pos)
    return match.group(0) if match else "?"


def _grammar_hint(segment: str, example: str, count: int) -> str:
    if segment not in WHITELIST_SEGMENTS:
        return f"{count} file(s) use invalid segment '{segment}', e.g. '{example}'"
    if re.match(r"^[a-z]+\d", example):
        return f"{count} file(s) missing the underscore, e.g. '{example}' (should be '{segment}_…')"
    if segment in SIDED_SEGMENTS and not example.endswith(("r", "v")):
        return f"{count} file(s) missing the r/v side, e.g. '{example}'"
    return f"{count} file(s) malformed for segment '{segment}', e.g. '{example}'"


def check_sequence(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    seqs = sorted(p.seq for p in parsed)
    findings = []
    if seqs[0] != 1:
        findings.append(Finding(item, f"{seqs[0]:04d}", "ERROR", "seq-start",
                                f"sequence starts at {seqs[0]:04d}, not 0001"))
    counts = defaultdict(int)
    for seq in seqs:
        counts[seq] += 1
    dups = [n for n, c in counts.items() if c > 1]
    if dups:
        findings.append(Finding(item, _compact(sorted(dups)), "ERROR", "seq-duplicate",
                                "sequence numbers used by more than one file"))
    gaps = [n for n in range(seqs[0], seqs[-1] + 1) if n not in counts]
    if gaps:
        findings.append(Finding(item, _compact(gaps), "ERROR", "seq-gap",
                                "gaps in the running sequence"))
    return findings


def check_duplicate_tokens(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    counts = defaultdict(list)
    for p in parsed:
        if p.pos:
            counts[p.pos].append(p.seq)
    return [
        Finding(item, _compact(sorted(seqs)), "ERROR", "duplicate-token",
                f"'{pos}' appears {len(seqs)} times")
        for pos, seqs in counts.items()
        if len(seqs) > 1
    ]


def check_segment_order(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    peak, peak_seg, findings = -1, None, []
    for p in sorted(parsed, key=lambda f: f.seq):
        family = p.position.family if p.position else None
        if family not in SEGMENT_RANK:
            continue
        rank = SEGMENT_RANK[family]
        if rank < peak:
            findings.append(Finding(item, f"{p.seq:04d}", "ERROR", "segment-order",
                                    f"'{p.pos}' comes after a '{peak_seg}' segment"))
        elif rank > peak:
            peak, peak_seg = rank, family
    return findings


def check_folio_order(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    findings = []
    for family in ("b", "f", "y"):
        run = sorted(
            (p.seq, p.position.number, p.position.side)
            for p in parsed
            if p.position and p.position.family == family and p.position.side
        )
        if run:
            findings.extend(_intra_leaf(item, family, run))
            findings.extend(_inter_leaf(item, family, run))
    return findings


def _intra_leaf(item: str, family: str, run: list[tuple[int, int, str]]) -> list[Finding]:
    first: dict[int, str] = {}
    for _, number, side in run:
        first.setdefault(number, side)
    r_first = sum(s == "r" for s in first.values())
    v_first = sum(s == "v" for s in first.values())
    if r_first and v_first:
        return [Finding(item, family, "ERROR", "rv-order",
                        f"'{family}' mixes recto-first and verso-first leaves")]
    if v_first:
        return [Finding(item, family, "REVIEW", "rv-order",
                        f"all '{family}' leaves run verso-before-recto")]
    return []


def _inter_leaf(item: str, family: str, run: list[tuple[int, int, str]]) -> list[Finding]:
    order = []
    for _, number, _ in run:
        if not order or order[-1] != number:
            order.append(number)
    # each leaf compared to the one before it, not the running max, so one
    # displaced leaf flags once at the discontinuity instead of on every leaf after it
    findings = [
        Finding(item, f"{family}{current}", "ERROR", "leaf-order",
                f"'{family}' numbering goes backward: {current} follows {previous}")
        for previous, current in zip(order, order[1:])
        if current < previous
    ]
    present = set(order)
    gaps = [n for n in range(min(present), max(present) + 1) if n not in present]
    if gaps:
        findings.append(Finding(item, _nums(gaps), "REVIEW", "leaf-gap",
                                f"'{family}' numbers skipped (confirm against the ms foliation): {_nums(gaps)}"))
    return findings


def check_inserts(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    seqs = sorted(p.seq for p in parsed if p.position and p.position.segment == "frag")
    if not seqs:
        return []
    return [
        Finding(item, _compact(list(range(start, end + 1))), "NOTE", "insert",
                f"insert present ({end - start + 1} image(s)) — review placement")
        for start, end in _runs(seqs)
    ]


def check_counters(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    findings = []
    pages = sorted(p.position.number for p in parsed
                   if p.position and p.position.family == "p")
    if pages and pages != list(range(1, len(pages) + 1)):
        findings.append(Finding(item, _nums(pages), "ERROR", "counter",
                                "'p' pagination should run contiguously from 1"))

    # a/z are fixed slots (z_01 inside back cover, z_02 back cover, ...), not a
    # running count, so a gap means a surface wasn't shot — REVIEW, not ERROR
    for segment in ("a", "z"):
        numbers = sorted(p.position.number for p in parsed
                         if p.position and p.position.family == segment)
        if not numbers:
            continue
        absent = [n for n in range(1, max(numbers)) if n not in numbers]
        if absent:
            findings.append(Finding(item, _nums(absent), "REVIEW", "slot-absent",
                                    f"'{segment}' slot(s) {_nums(absent)} not present — "
                                    "confirm those surfaces don't exist on the object"))
    return findings


def check_mixed_naming(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    # frag doesn't count toward "coded" — it's orthogonal to foliation status (see README)
    bare = sorted(p.seq for p in parsed if p.pos is None)
    coded = sum(1 for p in parsed
                if p.pos is not None and not (p.position and p.position.segment == "frag"))
    if not bare or not coded:
        return []
    return [Finding(item, _compact(bare), "ERROR", "mixed-naming",
                    f"{len(bare)} file(s) are bare sequence in an otherwise coded item "
                    f"({coded} coded) — foliation is all-or-nothing")]


def check_one_sided(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    sides = defaultdict(set)
    for p in parsed:
        pos = p.position
        if pos and pos.side and pos.family in SIDED_SEGMENTS:
            sides[(pos.family, pos.number)].add(pos.side)

    lonely = defaultdict(list)
    for (family, number), present in sides.items():
        if len(present) == 1:
            lonely[(family, next(iter(present)))].append(number)

    findings = []
    for (family, side), numbers in lonely.items():
        for start, end in _runs(sorted(numbers)):
            ref = f"{family}{start}" if start == end else f"{family}{start}-{family}{end}"
            noun = "leaf has" if start == end else f"{end - start + 1} leaves have"
            findings.append(Finding(item, ref, "REVIEW", "one-sided",
                                    f"{noun} only the {side} side named"))
    return findings


def check_frag_style(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    # frag_1r vs frag1r are both valid, but an item must pick one and stick to it
    styles = defaultdict(list)
    for p in parsed:
        if p.position and p.position.family == "frag" and p.position.canonical:
            style = "frag_" if p.pos.startswith("frag_") else "frag"
            styles[style].append((p.seq, p.pos))
    if len(styles) > 1:
        seqs = sorted(seq for entries in styles.values() for seq, _ in entries)
        examples = ", ".join(f"'{entries[0][1]}'" for entries in styles.values())
        return [Finding(item, _compact(seqs), "ERROR", "frag-style-mixed",
                        f"'frag' positions mix underscore and no-underscore styles in the same item, e.g. {examples}")]
    return []


def check_padding(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    findings = []
    for segment in ("a", "z"):
        widths = {len(p.position.raw_number) for p in parsed
                  if p.position and p.position.family == segment}
        if len(widths) > 1:
            findings.append(Finding(item, segment, "NOTE", "padding",
                                    f"'{segment}' positions mix zero-padding widths {sorted(widths)}"))
    return findings


def check_flyleaf_count(item: str, parsed: list[ParsedFile]) -> list[Finding]:
    for family in ("b", "y"):
        numbers = {p.position.number for p in parsed
                  if p.position and p.position.family == family}
        if len(numbers) > FLYLEAF_REVIEW_THRESHOLD:
            return [Finding(item, family, "REVIEW", "flyleaf-count",
                            f"{len(numbers)} '{family}' leaves — check whether any are "
                            "text-bearing preliminary/end leaves rather than blank flyleaves")]
    return []


def check_status(item: str, parsed: list[ParsedFile], status: str) -> list[Finding]:
    if status == "unfoliated":
        offenders = [p.seq for p in parsed if p.pos is not None]
        return _status_finding(item, offenders, "unfoliated-has-codes",
                               "carry position codes; unfoliated items are bare sequence")
    if status == "paginated":
        offenders = [p.seq for p in parsed
                     if p.position and (p.position.family == "f" or p.position.side)]
        return _status_finding(item, offenders, "paginated-not-p",
                               "use folio/recto-verso codes instead of p_")
    bare = [p.seq for p in parsed if p.pos is None]
    paged = [p.seq for p in parsed if p.position and p.position.family == "p"]
    return (_status_finding(item, bare, "foliated-bare", "have no position code on a foliated item")
            + _status_finding(item, paged, "foliated-has-p", "use p_ codes on a foliated item"))


def _status_finding(item: str, seqs: list[int], code: str, tail: str) -> list[Finding]:
    if not seqs:
        return []
    return [Finding(item, _compact(sorted(seqs)), "ERROR", code, f"{len(seqs)} file(s) {tail}")]


CHECKS = (
    check_grammar,
    check_sequence,
    check_duplicate_tokens,
    check_segment_order,
    check_folio_order,
    check_inserts,
    check_counters,
    check_mixed_naming,
    check_one_sided,
    check_padding,
    check_flyleaf_count,
    check_frag_style,
)


def run_checks(names: list[str], status: dict[str, str]) -> list[Finding]:
    by_item = defaultdict(list)
    for name in names:
        parsed = parse_filename(name)
        by_item[parsed.item if parsed else "UNPARSEABLE"].append((name, parsed))

    findings = []
    for item, rows in by_item.items():
        parsed = [p for _, p in rows if p]
        findings.extend(check_parsing(item, [n for n, _ in rows], parsed))
        if not parsed:
            continue
        for check in CHECKS:
            findings.extend(check(item, parsed))
        if item in status:
            findings.extend(check_status(item, parsed, status[item]))
    return findings


def _runs(numbers: list[int]) -> list[tuple[int, int]]:
    runs = []
    start = prev = numbers[0]
    for n in numbers[1:]:
        if n == prev + 1:
            prev = n
            continue
        runs.append((start, prev))
        start = prev = n
    runs.append((start, prev))
    return runs


def _format(numbers: list[int], pad: int) -> str:
    if not numbers:
        return ""
    pieces = []
    for start, end in _runs(numbers):
        lo, hi = f"{start:0{pad}d}", f"{end:0{pad}d}"
        pieces.append(lo if start == end else f"{lo}-{hi}")
    return ", ".join(pieces)


def _compact(numbers: list[int]) -> str:
    return _format(numbers, 4)


def _nums(numbers: list[int]) -> str:
    return _format(numbers, 1)


def _short(ref: str, limit: int = 40) -> str:
    if len(ref) <= limit:
        return ref
    parts = [p.strip() for p in ref.split(",")]
    return f"{parts[0]}…{parts[-1]}"


def print_report(source: str, names: list[str], findings: list[Finding]) -> None:
    rank = {"ERROR": 0, "REVIEW": 1, "NOTE": 2}
    items = sorted({parse_filename(n).item for n in names if parse_filename(n)})
    by_item = defaultdict(list)
    for finding in findings:
        by_item[finding.item].append(finding)

    tail = ["UNPARSEABLE"] if "UNPARSEABLE" in by_item else []
    print(f"\n{'=' * 70}\n{source}  —  {len(names)} files, {len(items)} items")
    for item in items + tail:
        hits = sorted(by_item.get(item, []), key=lambda f: (rank[f.severity], f.code))
        if not hits:
            print(f"\n  {item}: clean")
            continue
        print(f"\n  {item}:")
        for f in hits:
            print(f"    [{f.severity:<6}] {f.code:<16} {_short(f.ref):<18} {f.message}")

    totals = defaultdict(int)
    for finding in findings:
        totals[finding.severity] += 1
    line = ", ".join(f"{totals[s]} {s.lower()}" for s in rank if totals[s])
    print(f"\n  Totals: {line or 'nothing flagged'}")


def write_csv(path: Path, rows: list[tuple[str, Finding]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n", quoting=csv.QUOTE_ALL)
        writer.writerow(["manifest", "item", "reference", "severity", "code", "message"])
        for source, f in rows:
            writer.writerow([source, f.item, f.ref, f.severity, f.code, f.message])


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Check imaging manifest filenames against the naming convention.")
    parser.add_argument("manifests", nargs="+", type=Path)
    parser.add_argument("--status", type=Path, help="CSV of item,status (foliated/unfoliated/paginated)")
    parser.add_argument("--out", type=Path, help="combined CSV worklist path (default: output/<manifest>-report.csv beside this script)")
    args = parser.parse_args(argv)

    status = load_status(args.status) if args.status else {}
    csv_rows = []
    for manifest in args.manifests:
        names = load_manifest(manifest)
        findings = run_checks(names, status)
        print_report(manifest.name, names, findings)
        csv_rows.extend((manifest.name, f) for f in findings)

    out_path = args.out or OUTPUT_DIR / (
        f"{args.manifests[0].stem}-report.csv" if len(args.manifests) == 1 else "manifests-report.csv"
    )
    write_csv(out_path, csv_rows)
    print(f"\nWrote {len(csv_rows)} rows to {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
